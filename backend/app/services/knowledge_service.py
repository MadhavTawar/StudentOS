import hashlib
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, Dict, List, Optional, Tuple

import pymupdf
from docx import Document as DocxDocument
from fastapi import HTTPException, status
from sqlalchemy.orm import Session, joinedload

from backend.app.core.config import settings
from backend.app.core.logging import logger
from backend.app.models.knowledge import (
    KnowledgeAuditEvent,
    KnowledgeDocument,
    KnowledgeDocumentChunk,
    KnowledgeDocumentVersion,
)
from backend.app.models.user import User
from backend.app.models.workspace import WorkspaceMembership
from backend.app.schemas.knowledge import (
    DocumentVersionResponse,
    KnowledgeChatResponse,
    KnowledgeCitation,
    KnowledgeDocumentResponse,
    KnowledgeLibraryItem,
)
from backend.app.services.ai_service import AIService, sanitize_and_check_injection
from backend.app.services.knowledge_vector_store import KnowledgeVectorStore

PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
INSUFFICIENT_EVIDENCE_ANSWER = (
    "I couldn't find published workspace material that supports an answer to that question. "
    "Try a more specific question or ask a workspace owner."
)


def is_workspace_admin(db: Session, user: User, workspace_id: str) -> bool:
    return (
        db.query(WorkspaceMembership.id)
        .filter(
            WorkspaceMembership.user_id == user.id,
            WorkspaceMembership.workspace_id == workspace_id,
            WorkspaceMembership.role.in_(("owner", "admin")),
        )
        .first()
        is not None
    )


def can_manage_documents(db: Session, user: User, workspace_id: str) -> bool:
    return is_workspace_admin(db, user, workspace_id)


def require_document_manager(db: Session, user: User, workspace_id: str) -> None:
    if not can_manage_documents(db, user, workspace_id):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Workspace owner or admin access is required.",
        )


def _audit(
    db: Session,
    actor: User,
    action: str,
    entity_type: str,
    entity_id: str,
    details: Optional[Dict[str, Any]] = None,
) -> None:
    db.add(KnowledgeAuditEvent(
        actor_user_id=actor.id,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        details=json.dumps(details or {}, separators=(",", ":")),
    ))


def _resolve_storage_root() -> Path:
    root = Path(settings.KNOWLEDGE_STORAGE_DIR).expanduser().resolve()
    try:
        root.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        logger.exception("Could not open the configured private workspace document directory.")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Private workspace document storage is unavailable.",
        ) from exc
    return root


def _storage_path(storage_key: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}\.(pdf|docx)", storage_key):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="The stored document reference is invalid.",
        )
    root = _resolve_storage_root()
    path = (root / storage_key).resolve()
    if path.parent != root:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="The stored document path is invalid.",
        )
    return path


def _safe_filename(filename: str) -> str:
    name = PurePosixPath(filename.replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "workspace-document"


def _validate_and_identify_file(data: bytes, filename: str) -> Tuple[str, str]:
    if not data:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="The uploaded document is empty.",
        )
    if len(data) > settings.KNOWLEDGE_MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"Document exceeds the {settings.KNOWLEDGE_MAX_UPLOAD_BYTES}-byte upload limit.",
        )

    extension = Path(filename).suffix.lower()
    if extension == ".pdf" and b"%PDF-" in data[:1024]:
        return ".pdf", PDF_MEDIA_TYPE
    if extension != ".docx" or not data.startswith(b"PK"):
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Only valid PDF and DOCX documents are supported.",
        )

    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            expanded_size = sum(entry.file_size for entry in entries)
            if (
                expanded_size > settings.KNOWLEDGE_MAX_EXTRACTED_CHARS * 4
                or any(entry.file_size > 0 and entry.file_size / max(entry.compress_size, 1) > 100 for entry in entries)
                or "word/document.xml" not in archive.namelist()
                or any(entry.filename.lower().endswith("vbaproject.bin") for entry in entries)
            ):
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="The DOCX archive is invalid, too large after extraction, or contains macros.",
                )
            if archive.testzip() is not None:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="The DOCX archive failed integrity validation.",
                )
    except zipfile.BadZipFile as exc:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Only valid PDF and DOCX documents are supported.",
        ) from exc
    return ".docx", DOCX_MEDIA_TYPE


def _extract_pdf(data: bytes) -> Tuple[List[Tuple[Optional[int], str]], int]:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.is_encrypted:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Password-protected PDFs are not supported.",
                )
            if len(document) > settings.KNOWLEDGE_MAX_PAGES:
                raise HTTPException(
                    status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                    detail=f"PDF exceeds the {settings.KNOWLEDGE_MAX_PAGES}-page processing limit.",
                )
            pages: List[Tuple[Optional[int], str]] = []
            extracted_chars = 0
            for page_index in range(len(document)):
                page_number = page_index + 1
                page = document.load_page(page_index)
                page_text = page.get_text("text")
                if not isinstance(page_text, str):
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail="The PDF text extractor returned an unsupported page format.",
                    )
                text = page_text
                extracted_chars += len(text)
                if extracted_chars > settings.KNOWLEDGE_MAX_EXTRACTED_CHARS:
                    raise HTTPException(
                        status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                        detail="Document text exceeds the configured extraction limit.",
                    )
                if text.strip():
                    pages.append((page_number, text))
            return pages, len(document)
    except HTTPException:
        raise
    except (pymupdf.FileDataError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="The PDF could not be opened or its text could not be extracted.",
        ) from exc


def _extract_docx(data: bytes) -> Tuple[List[Tuple[Optional[int], str]], int]:
    try:
        document = DocxDocument(BytesIO(data))
    except (ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="The DOCX document could not be opened.",
        ) from exc

    content = [
        str(paragraph.text)
        for paragraph in document.paragraphs
        if str(paragraph.text).strip()
    ]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                str(cell.text).strip()
                for cell in row.cells
                if str(cell.text).strip()
            )
            if row_text:
                content.append(row_text)
    text = "\n".join(content)
    if len(text) > settings.KNOWLEDGE_MAX_EXTRACTED_CHARS:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail="Document text exceeds the configured extraction limit.",
        )
    return ([(None, text)] if text.strip() else []), 0


def _split_text(text: str, page: Optional[int], chunk_size: int, overlap: int) -> List[Dict[str, Any]]:
    if chunk_size < 200 or overlap < 0 or overlap >= chunk_size:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Institute document chunking settings are invalid.",
        )

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary
        content = text[start:end].strip()
        if content:
            chunks.append({
                "id": str(uuid.uuid4()),
                "ordinal": len(chunks),
                "page_start": page,
                "page_end": page,
                "content": content,
            })
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _make_chunks(pages: List[Tuple[Optional[int], str]]) -> List[Dict[str, Any]]:
    chunks: List[Dict[str, Any]] = []
    for page_number, text in pages:
        page_chunks = _split_text(
            text,
            page_number,
            settings.KNOWLEDGE_CHUNK_SIZE,
            settings.KNOWLEDGE_CHUNK_OVERLAP,
        )
        for chunk in page_chunks:
            chunk["ordinal"] = len(chunks)
            chunks.append(chunk)
    if not chunks:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No searchable text could be extracted. Scanned PDFs without a text layer are not supported.",
        )
    return chunks


def _mark_processing_failed(
    db: Session,
    actor: User,
    version_id: str,
    document_id: str,
    message: str,
) -> None:
    db.rollback()
    version = db.query(KnowledgeDocumentVersion).filter_by(id=version_id).first()
    if version is not None:
        version.status = "failed"
        version.error_message = message[:500]
        _audit(
            db,
            actor,
            "version_processing_failed",
            "document_version",
            version_id,
            {"reason": message[:120]},
        )
        db.commit()
    logger.warning("Institute document processing failed for document %s: %s", document_id, message[:120])


def upload_document(
    db: Session,
    actor: User,
    data: bytes,
    filename: str,
    *,
    workspace_id: str,
    title: Optional[str] = None,
    description: str = "",
    category: str = "General",
    document_id: Optional[str] = None,
) -> KnowledgeDocument:
    safe_filename = _safe_filename(filename)
    extension, media_type = _validate_and_identify_file(data, safe_filename)
    digest = hashlib.sha256(data).hexdigest()
    duplicate = db.query(KnowledgeDocumentVersion.id).filter(
        KnowledgeDocumentVersion.sha256 == digest,
        KnowledgeDocumentVersion.status.in_(("ready", "published", "superseded")),
        KnowledgeDocumentVersion.document.has(KnowledgeDocument.workspace_id == workspace_id),
    ).first()
    if duplicate is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="This document content has already been processed.",
        )

    if document_id:
        document = (
            db.query(KnowledgeDocument)
            .options(joinedload(KnowledgeDocument.versions))
            .filter(
                KnowledgeDocument.id == document_id,
                KnowledgeDocument.workspace_id == workspace_id,
            )
            .first()
        )
        if document is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")
        if document.status == "archived":
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Archived documents cannot be replaced.")
        version_number = max((version.version_number for version in document.versions), default=0) + 1
    else:
        normalized_title = (title or "").strip()
        if not normalized_title or len(normalized_title) > 240:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="A document title between 1 and 240 characters is required.",
            )
        normalized_category = category.strip() or "General"
        if len(normalized_category) > 100 or len(description) > 10000:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="Document category or description exceeds its allowed length.",
            )
        document = KnowledgeDocument(
            title=normalized_title,
            workspace_id=workspace_id,
            description=description.strip(),
            category=normalized_category,
            status="draft",
            uploaded_by=actor.id,
        )
        db.add(document)
        db.flush()
        version_number = 1

    document_pk = document.id
    storage_key = f"{uuid.uuid4().hex}{extension}"
    storage_path = _storage_path(storage_key)
    try:
        with storage_path.open("xb") as stored_file:
            stored_file.write(data)
    except OSError as exc:
        db.rollback()
        logger.exception("Could not store a workspace document.")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Private workspace document storage could not save the upload.",
        ) from exc

    version = KnowledgeDocumentVersion(
        document_id=document.id,
        version_number=version_number,
        original_filename=safe_filename,
        storage_key=storage_key,
        sha256=digest,
        media_type=media_type,
        status="processing",
        uploaded_by=actor.id,
    )
    db.add(version)
    db.flush()
    version_pk = version.id
    _audit(
        db,
        actor,
        "version_upload_started",
        "document_version",
        version_pk,
        {"version_number": version_number, "sha256": digest},
    )
    db.commit()

    try:
        pages, page_count = _extract_pdf(data) if extension == ".pdf" else _extract_docx(data)
        chunks = _make_chunks(pages)
        chunk_rows = [
            KnowledgeDocumentChunk(
                id=chunk["id"],
                version_id=version_pk,
                ordinal=chunk["ordinal"],
                page_start=chunk["page_start"],
                page_end=chunk["page_end"],
                content=chunk["content"],
            )
            for chunk in chunks
        ]
        db.add_all(chunk_rows)
        db.flush()
        KnowledgeVectorStore().upsert_chunks(
            chunks,
            document_id=document_pk,
            version_id=version_pk,
            workspace_id=workspace_id,
            status_value="ready",
        )
        version = db.query(KnowledgeDocumentVersion).filter_by(id=version_pk).first()
        if version is None:
            raise RuntimeError("The document version disappeared during processing.")
        version.page_count = page_count
        version.status = "ready"
        version.error_message = None
        version.processed_at = datetime.now(timezone.utc)
        _audit(
            db,
            actor,
            "version_processing_completed",
            "document_version",
            version_pk,
            {"chunk_count": len(chunks), "page_count": page_count},
        )
        db.commit()
    except HTTPException as exc:
        _mark_processing_failed(db, actor, version_pk, document_pk, exc.detail)
        raise
    except Exception as exc:
        _mark_processing_failed(db, actor, version_pk, document_pk, "Unexpected processing failure.")
        logger.exception("Unexpected workspace document processing failure.")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Institute document processing failed; this version remains unpublished.",
        ) from exc

    result = (
        db.query(KnowledgeDocument)
        .options(joinedload(KnowledgeDocument.versions))
        .filter(KnowledgeDocument.id == document_pk)
        .first()
    )
    if result is None:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Document record unavailable.")
    return result


def serialize_document(document: KnowledgeDocument) -> KnowledgeDocumentResponse:
    return KnowledgeDocumentResponse(
        id=document.id,
        title=document.title,
        description=document.description,
        category=document.category,
        status=document.status,
        created_at=document.created_at,
        updated_at=document.updated_at,
        versions=[
            DocumentVersionResponse(
                id=version.id,
                version_number=version.version_number,
                original_filename=version.original_filename,
                status=version.status,
                page_count=version.page_count,
                created_at=version.created_at,
                processed_at=version.processed_at,
                error_message=version.error_message,
            )
            for version in document.versions
        ],
    )


def list_documents(db: Session, workspace_id: str) -> List[KnowledgeDocument]:
    return (
        db.query(KnowledgeDocument)
        .filter(KnowledgeDocument.workspace_id == workspace_id)
        .options(joinedload(KnowledgeDocument.versions))
        .order_by(KnowledgeDocument.updated_at.desc())
        .all()
    )


def delete_workspace_documents(db: Session, workspace_id: str, actor: User) -> None:
    documents = (
        db.query(KnowledgeDocument)
        .filter(KnowledgeDocument.workspace_id == workspace_id)
        .options(joinedload(KnowledgeDocument.versions))
        .all()
    )
    versions = [version for document in documents for version in document.versions]
    if versions:
        vector_store = KnowledgeVectorStore()
        for version in versions:
            vector_store.delete_version(version.id)
        for version in versions:
            path = _storage_path(version.storage_key)
            try:
                path.unlink(missing_ok=True)
            except OSError as exc:
                logger.exception("Could not remove private workspace document file %s.", version.id)
                raise HTTPException(
                    status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                    detail="A private workspace document file could not be deleted.",
                ) from exc
    for document in documents:
        _audit(db, actor, "document_deleted", "knowledge_document", document.id, {"reason": "account_deletion"})
        db.delete(document)


def list_published_library(db: Session, workspace_id: str) -> List[KnowledgeLibraryItem]:
    rows = (
        db.query(KnowledgeDocument, KnowledgeDocumentVersion)
        .join(
            KnowledgeDocumentVersion,
            KnowledgeDocumentVersion.document_id == KnowledgeDocument.id,
        )
        .filter(
            KnowledgeDocument.status == "published",
            KnowledgeDocument.workspace_id == workspace_id,
            KnowledgeDocumentVersion.status == "published",
        )
        .order_by(KnowledgeDocument.category, KnowledgeDocument.title)
        .all()
    )
    return [
        KnowledgeLibraryItem(
            id=document.id,
            title=document.title,
            description=document.description,
            category=document.category,
            version_id=version.id,
            version_number=version.version_number,
            page_count=version.page_count,
            published_at=version.processed_at or version.created_at,
        )
        for document, version in rows
    ]


def publish_version(
    db: Session,
    actor: User,
    document_id: str,
    version_id: str,
    workspace_id: str,
) -> KnowledgeDocument:
    document = db.query(KnowledgeDocument).filter_by(id=document_id, workspace_id=workspace_id).first()
    version = db.query(KnowledgeDocumentVersion).filter_by(id=version_id, document_id=document_id).first()
    if document is None or version is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document version not found.")
    if version.status != "ready":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Only a successfully processed document version can be published.",
        )
    if db.query(KnowledgeDocumentChunk.id).filter_by(version_id=version.id).first() is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="This document version has no indexed chunks and cannot be published.",
        )

    previous_versions = db.query(KnowledgeDocumentVersion).filter(
        KnowledgeDocumentVersion.document_id == document.id,
        KnowledgeDocumentVersion.status == "published",
    ).all()
    vector_store = KnowledgeVectorStore()
    vector_store.set_version_status(version.id, "published")
    for previous in previous_versions:
        previous.status = "superseded"
    version.status = "published"
    document.status = "published"
    document.updated_at = datetime.now(timezone.utc)
    _audit(
        db,
        actor,
        "document_version_published",
        "document_version",
        version.id,
        {"document_id": document.id, "version_number": version.version_number},
    )
    db.commit()

    for previous in previous_versions:
        vector_store.set_version_status(previous.id, "superseded")
    return (
        db.query(KnowledgeDocument)
        .options(joinedload(KnowledgeDocument.versions))
        .filter_by(id=document.id)
        .first()
    )


def archive_document(db: Session, actor: User, document_id: str, workspace_id: str) -> KnowledgeDocument:
    document = db.query(KnowledgeDocument).filter_by(id=document_id, workspace_id=workspace_id).first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")
    versions = db.query(KnowledgeDocumentVersion).filter_by(document_id=document.id).all()
    was_archived = document.status == "archived"
    document.status = "archived"
    document.updated_at = datetime.now(timezone.utc)
    for version in versions:
        if version.status in {"published", "superseded", "ready"}:
            version.status = "archived"
    if was_archived is False:
        _audit(db, actor, "document_archived", "document", document.id)
    db.commit()

    vector_store = KnowledgeVectorStore()
    for version in versions:
        vector_store.set_version_status(version.id, "archived")
    return (
        db.query(KnowledgeDocument)
        .options(joinedload(KnowledgeDocument.versions))
        .filter_by(id=document.id)
        .first()
    )


def download_version(
    db: Session,
    actor: User,
    document_id: str,
    version_id: str,
    workspace_id: str,
) -> Tuple[Path, str]:
    row = (
        db.query(KnowledgeDocumentVersion, KnowledgeDocument)
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeDocumentVersion.document_id)
        .filter(
            KnowledgeDocument.id == document_id,
            KnowledgeDocument.workspace_id == workspace_id,
            KnowledgeDocumentVersion.id == version_id,
        )
        .first()
    )
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document version not found.")
    version, document = row
    if not can_manage_documents(db, actor, workspace_id) and not (
        document.status == "published" and version.status == "published"
    ):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document version not found.")
    path = _storage_path(version.storage_key)
    if not path.is_file():
        logger.error("Private document storage entry is missing for version %s.", version.id)
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="The requested document file is unavailable.",
        )
    _audit(db, actor, "document_downloaded", "document_version", version.id)
    db.commit()
    return path, version.original_filename


async def answer_question(db: Session, question: str, workspace_id: str) -> KnowledgeChatResponse:
    safe_question, suspicious = sanitize_and_check_injection(question.strip())
    if suspicious:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Instruction-like content is not accepted as a knowledge question.",
        )

    vector_results = KnowledgeVectorStore().search(safe_question, workspace_id=workspace_id, limit=24)
    candidate_ids = [item["id"] for item in vector_results]
    if not candidate_ids:
        return KnowledgeChatResponse(answer=INSUFFICIENT_EVIDENCE_ANSWER, insufficient_evidence=True)

    rows = (
        db.query(KnowledgeDocumentChunk, KnowledgeDocumentVersion, KnowledgeDocument)
        .join(KnowledgeDocumentVersion, KnowledgeDocumentVersion.id == KnowledgeDocumentChunk.version_id)
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeDocumentVersion.document_id)
        .filter(
            KnowledgeDocumentChunk.id.in_(candidate_ids),
            KnowledgeDocument.status == "published",
            KnowledgeDocument.workspace_id == workspace_id,
            KnowledgeDocumentVersion.status == "published",
        )
        .all()
    )
    allowed = {
        chunk.id: (chunk, version, document)
        for chunk, version, document in rows
    }
    max_distance = settings.KNOWLEDGE_MAX_DISTANCE
    citations: List[KnowledgeCitation] = []
    source_payload: List[Dict[str, Any]] = []
    for result in vector_results:
        row = allowed.get(result["id"])
        if row is None or result["distance"] is None or result["distance"] > max_distance:
            continue
        chunk, version, document = row
        metadata = result["metadata"]
        if metadata.get("document_id") != document.id or metadata.get("version_id") != version.id:
            continue
        _, suspicious_source = sanitize_and_check_injection(chunk.content)
        if suspicious_source:
            continue
        citations.append(KnowledgeCitation(
            chunk_id=chunk.id,
            document_id=document.id,
            document_title=document.title,
            version_id=version.id,
            version_number=version.version_number,
            page_start=chunk.page_start,
            page_end=chunk.page_end,
            excerpt=chunk.content[:700],
        ))
        source_payload.append({
            "source_id": chunk.id,
            "title": document.title,
            "page_start": chunk.page_start,
            "page_end": chunk.page_end,
            "content": chunk.content,
        })
        if len(citations) == 6:
            break

    if not source_payload:
        return KnowledgeChatResponse(answer=INSUFFICIENT_EVIDENCE_ANSWER, insufficient_evidence=True)

    answer = await AIService.get_provider().grounded_answer(safe_question, source_payload)
    return KnowledgeChatResponse(answer=answer, insufficient_evidence=False, citations=citations)
